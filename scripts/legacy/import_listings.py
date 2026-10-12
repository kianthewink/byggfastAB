"""Regenerate data/apartments.ts and data/lokaler.ts from an Excel workbook.

    npm run import:listings -- "C:\\path\\to\\bofast_listings_semicolon_arrays.xlsx"

Sheet "Listings" holds one row per listing. Array-typed fields are stored as
semicolon-separated strings (images, amenities, viewingDates, viewingTimes),
and `contacts` names an export from data/contacts.ts instead of inlining
people. Rows are routed to the two output files by `sourceCollection`.

Requires openpyxl (pip install openpyxl).
"""

from __future__ import annotations

import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required. Install it with:  pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent

COLLECTIONS = {
    "apartments": {
        "file": "data/apartments.ts",
        "export_name": "apartments",
        "doc": "/** Mock apartment listings used by the /lediga-lagenheter page. */",
    },
    "Lokaler": {
        "file": "data/lokaler.ts",
        "export_name": "Lokaler",
        "doc": "/** Commercial premises displayed on the /lediga-lokaler page. */",
    },
}

# Emitted in this order; columns not listed here are ignored.
SCALAR_FIELDS = [
    "id", "address", "city", "rooms", "area", "areaMax", "rent", "available",
    "floor", "floorsTotal", "elevator", "buildYear", "lat", "lng", "description",
]

NUMERIC = {"rooms", "area", "areaMax", "rent", "floor", "floorsTotal", "buildYear", "lat", "lng"}

TRUTHY = {"true", "ja", "yes", "1"}
FALSY = {"false", "nej", "no", "0"}


def split_list(value) -> list[str]:
    """"a;b;c" -> ["a", "b", "c"], dropping blanks."""
    if value is None:
        return []
    return [part.strip() for part in str(value).split(";") if part.strip()]


def clean(value):
    """Normalise a raw cell value; blank strings become None."""
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    return value


def to_number(value):
    try:
        num = float(value)
    except (TypeError, ValueError):
        return None
    return int(num) if num.is_integer() else num


def to_bool(value):
    if isinstance(value, bool):
        return value
    if value is None:
        return None
    text = str(value).strip().lower()
    if text in TRUTHY:
        return True
    if text in FALSY:
        return False
    return None


def quote(text: str) -> str:
    """Render a TS double-quoted string literal, keeping newlines as \\n."""
    escaped = (
        str(text)
        .replace("\\", "\\\\")
        .replace('"', '\\"')
        .replace("\r\n", "\\n")
        .replace("\n", "\\n")
    )
    return f'"{escaped}"'


def serialize(value, indent: int) -> str:
    pad = " " * (indent + 2)
    closing = " " * indent

    if isinstance(value, list):
        if not value:
            return "[]"
        items = ",\n".join(f"{pad}{serialize(v, indent + 2)}" for v in value)
        return f"[\n{items},\n{closing}]"

    if isinstance(value, dict):
        items = ",\n".join(
            f"{pad}{k}: {serialize(v, indent + 2)}" for k, v in value.items()
        )
        return f"{{\n{items},\n{closing}}}"

    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, str):
        return quote(value)
    return str(value)


def serialize_listing(listing: dict, contacts_ref: str | None) -> str:
    lines = [f"    {key}: {serialize(val, 4)}," for key, val in listing.items()]
    if contacts_ref:
        lines.append(f"    contacts: {contacts_ref},")
    body = "\n".join(lines)
    return f"  {{\n{body}\n  }}"


def main() -> None:
    if len(sys.argv) < 2:
        sys.exit("Usage: npm run import:listings -- <path-to-xlsx>")

    source = Path(sys.argv[1]).expanduser()
    if not source.is_file():
        sys.exit(f"File not found: {source}")

    workbook = openpyxl.load_workbook(source, data_only=True)
    sheet = workbook["Listings"] if "Listings" in workbook.sheetnames else workbook.worksheets[0]

    rows = sheet.iter_rows(values_only=True)
    headers = [clean(h) for h in next(rows)]

    buckets: dict[str, list[tuple[dict, str | None]]] = {k: [] for k in COLLECTIONS}
    warnings: list[str] = []

    for row_number, raw in enumerate(rows, start=2):
        row = {h: clean(v) for h, v in zip(headers, raw) if h}
        if not row.get("id"):
            continue

        collection = row.get("sourceCollection")
        if collection not in buckets:
            warnings.append(
                f'Row {row_number}: unknown sourceCollection "{collection}" - skipped.'
            )
            continue

        listing: dict = {}
        for field in SCALAR_FIELDS:
            value = row.get(field)
            if value is None:
                continue
            if field in NUMERIC:
                num = to_number(value)
                if num is not None:
                    listing[field] = num
            elif field == "elevator":
                flag = to_bool(value)
                if flag is not None:
                    listing[field] = flag
            else:
                listing[field] = value

        images = split_list(row.get("images"))
        if not images:
            warnings.append(f'Row {row_number} ({row["id"]}): no images.')
        listing["images"] = images

        amenities = split_list(row.get("amenities"))
        if amenities:
            listing["amenities"] = amenities

        # viewingDates / viewingTimes are parallel arrays zipped into Viewing objects.
        dates = split_list(row.get("viewingDates"))
        times = split_list(row.get("viewingTimes"))
        if len(dates) != len(times):
            warnings.append(
                f'Row {row_number} ({row["id"]}): {len(dates)} viewingDates vs '
                f"{len(times)} viewingTimes - extras dropped."
            )
        viewings = [{"date": d, "time": t} for d, t in zip(dates, times)]
        if viewings:
            listing["viewings"] = viewings

        buckets[collection].append((listing, row.get("contacts")))

    # Excel has no unique-key constraint, but ids become URLs - so enforce it here.
    for collection, entries in buckets.items():
        seen: dict[str, int] = {}
        for listing, _ in entries:
            listing_id = listing["id"]
            count = seen.get(listing_id, 0)
            seen[listing_id] = count + 1
            if count:
                unique = f"{listing_id}-{count + 1}"
                warnings.append(
                    f'{collection}: duplicate id "{listing_id}" - renamed to '
                    f'"{unique}" to keep routes unique.'
                )
                listing["id"] = unique

    for collection, config in COLLECTIONS.items():
        entries = buckets[collection]
        refs = sorted({ref for _, ref in entries if ref})

        body = ",\n".join(serialize_listing(listing, ref) for listing, ref in entries)
        import_line = (
            f'import {{ {", ".join(refs)} }} from "@/data/contacts";\n' if refs else ""
        )

        contents = (
            "// AUTO-GENERATED by scripts/import_listings.py - do not edit by hand.\n"
            "// Re-run: npm run import:listings -- <path-to-xlsx>\n\n"
            'import type { Apartment } from "@/lib/types";\n'
            f"{import_line}"
            f"\n{config['doc']}\n"
            f"export const {config['export_name']}: Apartment[] = [\n{body},\n];\n"
        )

        (ROOT / config["file"]).write_text(contents, encoding="utf-8")
        print(f"OK {config['file']} - {len(entries)} listings")

    if warnings:
        print("\nWarnings:")
        for warning in warnings:
            print(f"  ! {warning}")


if __name__ == "__main__":
    main()
